"""The public demo: daily limits per address and overall, export without a second
extraction, CORS for the web page, the Lambda warm-up."""

import io
import json
from typing import Any

import httpx
import pytest
from api_support import API_KEY, SCHEMA, FakePipeline, png
from openpyxl import load_workbook

from schemalens_api.deps import get_pipeline
from schemalens_api.limits import DynamoLimiter, MemoryLimiter, client_key
from schemalens_api.main import create_app
from schemalens_core.errors import LlmFailed

ORIGIN = "https://schemalens.space"
DEMO_KEY = "demo-key"


def upload(document: bytes | None = None, **fields: Any) -> dict[str, Any]:
    data = {k: json.dumps(v) if isinstance(v, dict) else v for k, v in fields.items()}
    return {"files": {"file": ("doc.png", document or png(), "image/png")}, "data": data}


@pytest.fixture
def demo(make_settings):
    """An app with small limits and a fake pipeline; the client sends the demo key."""

    async def build(**overrides: Any):
        settings = make_settings(
            **{
                "demo_daily_documents": 2,
                "allowed_origins": ORIGIN,
                "demo_keys": DEMO_KEY,
                **overrides,
            }
        )
        app = create_app(settings)
        fake = FakePipeline(settings)
        app.dependency_overrides[get_pipeline] = lambda: fake
        client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://t",
            headers={"Authorization": f"Bearer {DEMO_KEY}"},
        )
        return client, fake

    return build


async def test_each_address_gets_a_few_documents_a_day(demo) -> None:
    client, fake = await demo()
    statuses = [
        (await client.post("/api/v1/extract", **upload(schema=SCHEMA))).status_code
        for _ in range(3)
    ]
    assert statuses == [200, 200, 429]
    assert len(fake.calls) == 2
    response = await client.post("/api/v1/extract", **upload(schema=SCHEMA))
    assert response.headers["retry-after"] and response.json()["type"].endswith("limit-reached")


async def test_a_service_key_is_not_limited_per_address(demo) -> None:
    client, fake = await demo(demo_daily_documents=1)
    service = {"Authorization": f"Bearer {API_KEY}"}
    statuses = [
        (await client.post("/api/v1/extract", headers=service, **upload(schema=SCHEMA))).status_code
        for _ in range(3)
    ]
    assert statuses == [200, 200, 200] and len(fake.calls) == 3


async def test_the_whole_demo_has_a_daily_ceiling(demo) -> None:
    client, _ = await demo(demo_daily_documents=10, demo_global_daily_documents=1)
    assert (await client.post("/api/v1/extract", **upload(schema=SCHEMA))).status_code == 200
    assert (await client.post("/api/v1/extract", **upload(schema=SCHEMA))).status_code == 429


async def test_rejected_requests_do_not_use_the_limit(demo) -> None:
    client, _ = await demo(demo_daily_documents=1)
    bad = await client.post("/api/v1/extract", **upload(schema={"type": "object"}))
    assert bad.status_code == 422
    assert (await client.post("/api/v1/extract", **upload(schema=SCHEMA))).status_code == 200


def events(response: httpx.Response) -> list[dict[str, Any]]:
    return [json.loads(line) for line in response.text.splitlines()]


async def test_stream_reports_each_stage_then_the_result(demo) -> None:
    client, _ = await demo()
    response = await client.post("/api/v1/extract?stream=true", **upload(schema=SCHEMA))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/x-ndjson")
    lines = events(response)
    assert [line["event"] for line in lines] == [
        "accepted",
        "reading",
        "read",
        "extracting",
        "result",
    ]
    assert lines[0]["pages"] == 1 and lines[0]["bytes"] > 0
    assert lines[2]["lines"] == 12 and lines[3]["fields"] == 3
    assert lines[-1]["result"]["data"]["sklep"] == "SKLEP EXAMPLE"


async def test_stream_failure_is_the_last_event_with_the_usual_problem(demo) -> None:
    client, fake = await demo()
    fake.error = LlmFailed("the model answered nonsense")
    lines = events(await client.post("/api/v1/extract?stream=true", **upload(schema=SCHEMA)))
    assert lines[-1]["event"] == "error"
    assert lines[-1]["problem"]["status"] == 502
    assert lines[-1]["problem"]["type"].endswith("llm-failed")


async def test_stream_checks_the_request_and_the_limit_before_streaming(demo) -> None:
    client, _ = await demo(demo_daily_documents=1)
    bad = await client.post("/api/v1/extract?stream=true", **upload(schema={"type": "object"}))
    assert bad.status_code == 422
    assert (
        await client.post("/api/v1/extract?stream=true", **upload(schema=SCHEMA))
    ).status_code == 200
    assert (
        await client.post("/api/v1/extract?stream=true", **upload(schema=SCHEMA))
    ).status_code == 429


async def test_export_turns_a_result_into_a_file_without_extracting_again(demo) -> None:
    client, fake = await demo()
    result = (await client.post("/api/v1/extract", **upload(schema=SCHEMA))).json()
    calls = len(fake.calls)

    csv = await client.post("/api/v1/render?format=csv", json={"result": result, "schema": SCHEMA})
    xlsx = await client.post(
        "/api/v1/render?format=xlsx", json={"result": result, "schema": SCHEMA}
    )
    as_json = await client.post(
        "/api/v1/render?format=json", json={"result": result, "schema": SCHEMA}
    )

    assert csv.content.decode("utf-8-sig").splitlines()[0] == "sklep,suma"
    assert load_workbook(io.BytesIO(xlsx.content)).sheetnames == ["data", "meta"]
    assert as_json.json()["data"] == result["data"]
    assert "attachment" in as_json.headers["content-disposition"]
    assert len(fake.calls) == calls  # nothing extracted again


async def test_demo_info_tells_the_site_its_limits(demo) -> None:
    client, _ = await demo()
    info = (await client.get("/api/v1/demo")).json()
    assert info == {
        "max_upload_bytes": 4 * 1024 * 1024,
        "max_pages": 3,
        "daily_documents": 2,
    }


async def test_cors_lets_only_the_landing_site_in(demo) -> None:
    client, _ = await demo()
    headers = {"Origin": ORIGIN, "Access-Control-Request-Method": "POST"}
    allowed = await client.options("/api/v1/extract", headers=headers)
    other = await client.options(
        "/api/v1/extract", headers={**headers, "Origin": "https://evil.example"}
    )
    assert allowed.headers["access-control-allow-origin"] == ORIGIN
    assert "access-control-allow-origin" not in other.headers


# --- API keys ---------------------------------------------------------------------------


async def test_every_call_needs_a_key(demo) -> None:
    client, fake = await demo(api_keys="k-one, k-two")
    client.headers.pop("Authorization")
    assert (await client.get("/healthz")).status_code == 200  # health and docs stay open
    assert (await client.get("/docs")).status_code == 200
    missing = await client.post("/api/v1/extract", **upload(schema=SCHEMA))
    assert missing.status_code == 401 and missing.headers["www-authenticate"] == "Bearer"
    wrong = await client.get("/api/v1/presets", headers={"Authorization": "Bearer nope"})
    assert wrong.status_code == 401
    bearer = await client.get("/api/v1/presets", headers={"Authorization": "Bearer k-two"})
    header = await client.get("/api/v1/presets", headers={"X-API-Key": "k-one"})
    assert bearer.status_code == header.status_code == 200
    assert (await client.get("/api/v1/demo", headers={"X-API-Key": DEMO_KEY})).status_code == 200
    assert fake.calls == []


async def test_the_api_refuses_to_start_without_any_key(make_settings) -> None:
    app = create_app(make_settings(api_keys=""))
    with pytest.raises(RuntimeError, match="keys new"):
        async with app.router.lifespan_context(app):
            pass
    for configured in ({"demo_keys": DEMO_KEY}, {"api_keys": "k"}):
        app = create_app(make_settings(**{"api_keys": "", **configured}))
        async with app.router.lifespan_context(app):
            pass


GATEWAY = {"x-amzn-request-context": json.dumps({"identity": {"apiKeyId": "abc123"}})}


async def test_on_aws_a_key_api_gateway_checked_is_enough(make_settings) -> None:
    settings = make_settings(api_keys="", gateway_keys=True, lambda_request_context=True)
    app = create_app(settings)
    app.dependency_overrides[get_pipeline] = lambda: FakePipeline(settings)
    client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")
    assert (await client.get("/api/v1/presets", headers=GATEWAY)).status_code == 200
    # The demo's Function URL: a request context without a key id, whatever the client sends.
    function_url = {"x-amzn-request-context": json.dumps({"http": {"sourceIp": "203.0.113.7"}})}
    assert (await client.get("/api/v1/presets", headers=function_url)).status_code == 401
    assert (await client.get("/api/v1/presets", headers={"X-API-Key": "abc123"})).status_code == 401


async def test_the_gateway_context_counts_only_on_aws(make_settings) -> None:
    app = create_app(make_settings(gateway_keys=True))  # lambda_request_context off
    client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")
    assert (await client.get("/api/v1/presets", headers=GATEWAY)).status_code == 401


async def test_a_key_is_limited_per_key_not_per_address(demo) -> None:
    client, _ = await demo(api_keys="k-one,k-two", key_daily_documents=1, demo_daily_documents=10)
    one = {"Authorization": "Bearer k-one"}
    two = {"Authorization": "Bearer k-two"}
    assert (
        await client.post("/api/v1/extract", headers=one, **upload(schema=SCHEMA))
    ).status_code == 200
    assert (
        await client.post("/api/v1/extract", headers=one, **upload(schema=SCHEMA))
    ).status_code == 429
    assert (
        await client.post("/api/v1/extract", headers=two, **upload(schema=SCHEMA))
    ).status_code == 200


# --- limiters ---------------------------------------------------------------------------


def test_addresses_are_hashed_with_a_salt() -> None:
    key = client_key("extract", "203.0.113.7", "salt")
    assert "203.0.113.7" not in key and key != client_key("extract", "203.0.113.7", "other")


def test_memory_limiter_resets_every_day() -> None:
    now = [1_800_000_000.0]
    limiter = MemoryLimiter(clock=lambda: now[0])
    assert limiter.take("k", 1) and not limiter.take("k", 1)
    now[0] += 24 * 3600
    assert limiter.take("k", 1)


class FakeDynamo:
    """Mimics the conditional ADD the DynamoDB limiter relies on."""

    def __init__(self) -> None:
        self.items: dict[str, dict[str, int]] = {}

    def update_item(self, **request: Any) -> None:
        from botocore.exceptions import ClientError

        key = request["Key"]["pk"]["S"]
        limit = int(request["ExpressionAttributeValues"][":limit"]["N"])
        item = self.items.setdefault(key, {})
        if item.get("uses", 0) >= limit:
            raise ClientError({"Error": {"Code": "ConditionalCheckFailedException"}}, "UpdateItem")
        item["uses"] = item.get("uses", 0) + 1
        item.setdefault("expires_at", int(request["ExpressionAttributeValues"][":ttl"]["N"]))


def test_dynamo_limiter_counts_atomically_with_a_ttl(make_settings) -> None:
    table = FakeDynamo()
    limiter = DynamoLimiter(
        make_settings(limits_table="limits"), client=table, clock=lambda: 1_800_000_000.0
    )
    assert [limiter.take("extract:abc", 2) for _ in range(3)] == [True, True, False]
    (item,) = table.items.values()
    assert item == {"uses": 2, "expires_at": 1_800_000_000 + 2 * 24 * 3600}


async def test_warm_up_ping_loads_the_models_without_a_document(make_settings) -> None:
    warmed: list[bool] = []

    class Pipeline:
        def warm_up(self) -> None:
            warmed.append(True)

    app = create_app(make_settings())
    app.state.pipeline = Pipeline()
    client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")
    response = await client.post("/internal/warm-up", json={"warmup": True})
    assert response.json() == {"warm": True}
    assert warmed == [True]
    assert "/internal/warm-up" not in app.openapi()["paths"]


def _counted_address(demo_app, headers: dict[str, str]) -> str:
    from starlette.requests import Request

    from schemalens_api.deps import client_address

    scope = {
        "type": "http",
        "app": demo_app,
        "headers": [(k.lower().encode(), v.encode()) for k, v in headers.items()],
        "client": ("127.0.0.1", 5000),
    }
    return client_address(Request(scope))


CONTEXT = {"x-amzn-request-context": json.dumps({"http": {"sourceIp": "203.0.113.7"}})}


def test_on_aws_the_address_comes_from_the_invocation_context(make_settings) -> None:
    app = create_app(make_settings(lambda_request_context=True))
    assert _counted_address(app, CONTEXT) == "203.0.113.7"
    assert _counted_address(app, {}) == "unknown"  # never the adapter's own 127.0.0.1


def test_elsewhere_the_context_header_is_ignored(make_settings) -> None:
    app = create_app(make_settings())
    assert _counted_address(app, CONTEXT) == "127.0.0.1"
