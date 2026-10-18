import csv
import io
from datetime import date, datetime
from typing import Any

import httpx
import pytest
from api_support import AUTH, FakePipeline, png
from openpyxl import load_workbook

from schemalens_api.formats import MEDIA_TYPES, negotiate, to_csv, to_xlsx
from schemalens_api.problems import ApiError

XLSX = MEDIA_TYPES["xlsx"]
STR, NUM = {"type": "string"}, {"type": "number"}
ITEM = {"type": "object", "properties": {"nazwa": STR, "ilosc": {"type": "integer"}, "cena": NUM}}

FLAT = {
    "type": "object",
    "properties": {
        "numer": STR,
        "data": {"type": "string", "format": "date"},
        "kwota": NUM,
        "zaplacono": {"type": "boolean"},
    },
}
WITH_OBJECT = {
    "type": "object",
    "properties": {
        "numer": STR,
        "seller": {
            "type": "object",
            "properties": {
                "name": STR,
                "nip": STR,
                "address": {"type": "object", "properties": {"city": STR}},
            },
        },
    },
}
ONE_LIST = {
    "type": "object",
    "properties": {
        "numer": STR,
        "pozycje": {"type": "array", "items": ITEM},
        "kwota": NUM,
        "tagi": {"type": "array", "items": STR},
    },
}
TWO_LISTS = {
    "type": "object",
    "properties": {
        "numer": STR,
        "pozycje": {"type": "array", "items": ITEM},
        "platnosci": {
            "type": "array",
            "items": {"type": "object", "properties": {"data": STR, "kwota": NUM}},
        },
    },
}
POZYCJE = [
    {"nazwa": "Paliwo", "ilosc": 2, "cena": 150.0},
    {"nazwa": "Myjnia", "ilosc": 1, "cena": 12.45},
]


def body(data: dict[str, Any], warnings: list[str] | None = None) -> dict[str, Any]:
    return {
        "data": data,
        "fields": {
            name: {"confidence": 0.9 if value is not None else 0.0, "evidence": ["L001", "L002"]}
            for name, value in data.items()
        },
        "warnings": warnings or [],
        "usage": {"pages": 1, "ocr_ms": 120, "llm_ms": 340, "model": "m"},
    }


def rows(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM, so Excel reads Polish characters
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


# --- negotiation ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("param", "accept", "expected"),
    [
        (None, None, "json"),
        (None, "", "json"),
        (None, "*/*", "json"),
        (None, "application/json", "json"),
        (None, "text/csv", "csv"),
        (None, XLSX, "xlsx"),
        (None, "text/html, text/csv;q=0.9, */*;q=0.1", "csv"),
        ("csv", "application/json", "csv"),  # the query parameter wins
        ("xlsx", None, "xlsx"),
        ("json", "text/csv", "json"),
    ],
)
def test_negotiate(param: str | None, accept: str | None, expected: str) -> None:
    assert negotiate(param, accept) == expected


@pytest.mark.parametrize(
    ("param", "accept"),
    [
        ("pdf", None),
        ("", None),
        ("CSV", None),
        (None, "application/xml"),
        (None, "image/png, text/html"),
    ],
)
def test_negotiate_rejects_unknown_formats(param: str | None, accept: str | None) -> None:
    with pytest.raises(ApiError) as excinfo:
        negotiate(param, accept)
    assert excinfo.value.status == 406 and "format=json|csv|xlsx" in excinfo.value.detail


# --- CSV --------------------------------------------------------------------------------


def test_csv_flat_schema() -> None:
    content, partial = to_csv(
        body({"numer": "FV/1", "data": "2026-10-03", "kwota": 1230.5, "zaplacono": True}), FLAT
    )

    assert rows(content) == [
        ["numer", "data", "kwota", "zaplacono"],
        ["FV/1", "2026-10-03", "1230.5", "true"],
    ]
    assert partial is False
    assert content.endswith(b"\r\n") and b";" not in content


def test_csv_null_values_are_empty_cells() -> None:
    content, _ = to_csv(
        body({"numer": None, "data": None, "kwota": None, "zaplacono": False}), FLAT
    )
    assert rows(content)[1] == ["", "", "", "false"]


def test_csv_objects_are_flattened_with_dotted_paths() -> None:
    data = {
        "numer": "FV/1",
        "seller": {"name": "Example Sp. z o.o.", "nip": "1234563218", "address": {"city": "Łódź"}},
    }

    content, partial = to_csv(body(data), WITH_OBJECT)

    assert rows(content) == [
        ["numer", "seller.name", "seller.nip", "seller.address.city"],
        ["FV/1", "Example Sp. z o.o.", "1234563218", "Łódź"],
    ]
    assert partial is False
    assert rows(to_csv(body({"numer": "FV/1", "seller": None}), WITH_OBJECT)[0])[1] == [
        "FV/1",
        "",
        "",
        "",
    ]


def test_csv_one_list_becomes_one_row_per_item_with_scalars_repeated() -> None:
    data = {"numer": "FV/1", "pozycje": POZYCJE, "kwota": 312.45, "tagi": ["paliwo", "auto"]}

    content, partial = to_csv(body(data), ONE_LIST)

    assert rows(content) == [
        ["numer", "pozycje.nazwa", "pozycje.ilosc", "pozycje.cena", "kwota", "tagi"],
        ["FV/1", "Paliwo", "2", "150.0", "312.45", "paliwo | auto"],
        ["FV/1", "Myjnia", "1", "12.45", "312.45", "paliwo | auto"],
    ]
    assert partial is False


@pytest.mark.parametrize("empty", [[], None])
def test_csv_empty_list_gives_one_row_with_empty_list_columns(empty: Any) -> None:
    content, _ = to_csv(
        body({"numer": "FV/1", "pozycje": empty, "kwota": 5.0, "tagi": []}), ONE_LIST
    )

    assert rows(content) == [
        ["numer", "pozycje.nazwa", "pozycje.ilosc", "pozycje.cena", "kwota", "tagi"],
        ["FV/1", "", "", "", "5.0", ""],
    ]


def test_csv_second_list_is_json_text_and_flattening_is_partial() -> None:
    data = {
        "numer": "FV/1",
        "pozycje": POZYCJE,
        "platnosci": [{"data": "2026-10-03", "kwota": 100.0}],
    }

    content, partial = to_csv(body(data), TWO_LISTS)

    table = rows(content)
    assert table[0] == ["numer", "pozycje.nazwa", "pozycje.ilosc", "pozycje.cena", "platnosci"]
    assert len(table) == 3
    assert table[1][4] == table[2][4] == '[{"data": "2026-10-03", "kwota": 100.0}]'
    assert partial is True


def test_csv_lists_nested_in_items_or_objects_are_json_text_and_partial() -> None:
    nested_item = {
        "type": "object",
        "properties": {
            "nazwa": STR,
            "rabaty": {"type": "array", "items": {"type": "object", "properties": {"kod": STR}}},
        },
    }
    in_items = {
        "type": "object",
        "properties": {"pozycje": {"type": "array", "items": nested_item}},
    }
    in_object = {
        "type": "object",
        "properties": {
            "seller": {
                "type": "object",
                "properties": {"name": STR, "kontakty": {"type": "array", "items": ITEM}},
            }
        },
    }

    first, partial_items = to_csv(
        body({"pozycje": [{"nazwa": "A", "rabaty": [{"kod": "X"}]}]}), in_items
    )
    second, partial_object = to_csv(
        body({"seller": {"name": "E", "kontakty": [{"nazwa": "N"}]}}), in_object
    )

    assert rows(first) == [["pozycje.nazwa", "pozycje.rabaty"], ["A", '[{"kod": "X"}]']]
    assert rows(second) == [["seller.name", "seller.kontakty"], ["E", '[{"nazwa": "N"}]']]
    assert partial_items is True and partial_object is True


def test_csv_include_meta_adds_confidence_for_scalar_fields_only() -> None:
    data = {"numer": "FV/1", "pozycje": POZYCJE[:1], "kwota": None, "tagi": ["a"]}

    content, _ = to_csv(body(data), ONE_LIST, include_meta=True)

    assert rows(content) == [
        [
            "numer",
            "numer__confidence",
            "pozycje.nazwa",
            "pozycje.ilosc",
            "pozycje.cena",
            "kwota",
            "kwota__confidence",
            "tagi",
        ],
        ["FV/1", "0.9", "Paliwo", "2", "150.0", "", "0.0", "a"],
    ]


def test_csv_quotes_commas_and_defuses_formulas() -> None:
    data = {
        "numer": '=HYPERLINK("http://evil.example","x")',
        "data": "+48 600",
        "kwota": -5.5,
        "zaplacono": None,
    }

    table = rows(to_csv(body(data), FLAT)[0])

    assert table[1][0] == '\'=HYPERLINK("http://evil.example","x")'
    assert table[1][1] == "'+48 600"
    assert table[1][2] == "-5.5"  # a real negative number stays a number
    assert (
        rows(
            to_csv(body({"numer": "a, b", "data": None, "kwota": None, "zaplacono": None}), FLAT)[0]
        )[1][0]
        == "a, b"
    )


# --- XLSX -------------------------------------------------------------------------------


def workbook(content: bytes):
    return load_workbook(io.BytesIO(content))


def test_xlsx_round_trip_flat_schema_with_typed_cells() -> None:
    data = {"numer": "FV/1", "data": "2026-10-03", "kwota": 1230.5, "zaplacono": True}

    book = workbook(
        to_xlsx(
            body(
                data,
                ["numer: possible OCR noise", "document truncated to fit context (2 of 4 pages)"],
            ),
            FLAT,
        )
    )

    assert book.sheetnames == ["data", "meta"]
    sheet = book["data"]
    assert [cell.value for cell in sheet[1]] == ["numer", "data", "kwota", "zaplacono"]
    numer, when, amount, paid = sheet[2]
    assert numer.value == "FV/1"
    assert (
        isinstance(when.value, datetime) and when.value.date() == date(2026, 10, 3) and when.is_date
    )
    assert amount.value == 1230.5 and amount.data_type == "n"
    assert paid.value is True

    meta = [[cell.value for cell in row] for row in book["meta"].iter_rows()]
    assert meta[0] == ["field", "confidence", "evidence", "warnings"]
    assert meta[1] == ["numer", 0.9, "L001, L002", "numer: possible OCR noise"]
    assert ["(document)", None, None, "document truncated to fit context (2 of 4 pages)"] in meta
    assert ["usage", "value"] in [row[:2] for row in meta]
    assert ["ocr_ms", 120] in [row[:2] for row in meta] and ["model", "m"] in [
        row[:2] for row in meta
    ]


def test_xlsx_is_styled_for_reading_with_the_layout_unchanged() -> None:
    data = {"numer": "FV/1", "data": "2026-10-03", "kwota": 1230.5, "zaplacono": True}
    book = workbook(to_xlsx(body(data), FLAT))
    sheet = book["data"]
    header = sheet["A1"]
    assert header.font.bold and header.fill.start_color.rgb.endswith("FFE81A")
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:D2"
    assert sheet["C2"].number_format == "#,##0.00"
    assert sheet.column_dimensions["A"].width >= 10
    confidence = book["meta"]["B2"]
    assert confidence.number_format == "0%" and confidence.fill.start_color.rgb.endswith("D9F2D9")


def test_xlsx_headings_use_the_labels_when_given() -> None:
    data = {"numer": "FV/1", "data": None, "kwota": None, "zaplacono": None}
    book = workbook(to_xlsx(body(data), FLAT, {"numer": "Numer faktury", "kwota": "Kwota"}))
    assert [c.value for c in book["data"][1]] == ["Numer faktury", "Data", "Kwota", "Zaplacono"]
    assert book["meta"]["A2"].value == "Numer faktury"


def test_xlsx_objects_stay_on_the_data_sheet() -> None:
    data = {
        "numer": "FV/1",
        "seller": {"name": "Example", "nip": "1234563218", "address": {"city": "Łódź"}},
    }

    book = workbook(to_xlsx(body(data), WITH_OBJECT))

    assert book.sheetnames == ["data", "meta"]
    assert [c.value for c in book["data"][1]] == [
        "numer",
        "seller.name",
        "seller.nip",
        "seller.address.city",
    ]
    assert [c.value for c in book["data"][2]] == ["FV/1", "Example", "1234563218", "Łódź"]


def test_xlsx_every_list_of_objects_gets_its_own_sheet() -> None:
    data = {
        "numer": "FV/1",
        "pozycje": POZYCJE,
        "platnosci": [{"data": "2026-10-03", "kwota": 100.0}],
    }

    book = workbook(to_xlsx(body(data), TWO_LISTS))

    assert book.sheetnames == ["data", "pozycje", "platnosci", "meta"]
    assert [[c.value for c in row] for row in book["data"].iter_rows()] == [["numer"], ["FV/1"]]
    assert [[c.value for c in row] for row in book["pozycje"].iter_rows()] == [
        ["nazwa", "ilosc", "cena"],
        ["Paliwo", 2, 150.0],
        ["Myjnia", 1, 12.45],
    ]
    assert [[c.value for c in row] for row in book["platnosci"].iter_rows()] == [
        ["data", "kwota"],
        ["2026-10-03", 100.0],
    ]


def test_xlsx_empty_list_sheet_has_only_its_header_and_scalar_lists_are_joined() -> None:
    book = workbook(
        to_xlsx(body({"numer": "FV/1", "pozycje": [], "kwota": 1.0, "tagi": ["a", "b"]}), ONE_LIST)
    )

    assert [[c.value for c in row] for row in book["pozycje"].iter_rows()] == [
        ["nazwa", "ilosc", "cena"]
    ]
    assert [c.value for c in book["data"][2]] == ["FV/1", 1.0, "a | b"]


def test_xlsx_sheet_names_are_trimmed_safe_and_unique() -> None:
    long_name = "bardzo_dluga_nazwa_listy_pozycji_faktury_korygujacej"
    schema = {
        "type": "object",
        "properties": {
            long_name: {"type": "array", "items": ITEM},
            long_name + "_2": {"type": "array", "items": ITEM},
            "meta": {"type": "array", "items": ITEM},
            "a/b": {"type": "object", "properties": {"lista": {"type": "array", "items": ITEM}}},
        },
    }

    book = workbook(to_xlsx(body(dict.fromkeys(schema["properties"])), schema))

    names = book.sheetnames
    assert names[0] == "data" and names[-1] == "meta" and len(set(names)) == len(names) == 6
    assert all(len(name) <= 31 for name in names)
    assert names[1] == long_name[:31] and names[2] == long_name[:29] + "_2"
    assert "meta_2" in names and "a_b.lista" in names


def test_xlsx_text_is_never_stored_as_a_formula() -> None:
    book = workbook(
        to_xlsx(body({"numer": "=1+1", "data": None, "kwota": None, "zaplacono": None}), FLAT)
    )
    cell = book["data"]["A2"]
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_xlsx_bad_date_text_stays_text() -> None:
    book = workbook(
        to_xlsx(body({"numer": "x", "data": "03.10.2026", "kwota": None, "zaplacono": None}), FLAT)
    )
    assert book["data"]["B2"].value == "03.10.2026"


# --- over HTTP --------------------------------------------------------------------------

INVOICE_DATA = {
    "numer": "FV/1",
    "pozycje": POZYCJE,
    "platnosci": [{"data": "2026-10-03", "kwota": 100.0}],
}


def upload(schema: dict[str, Any]) -> dict[str, Any]:
    import json

    return {
        "files": {"file": ("doc.png", png(), "image/png")},
        "data": {"schema": json.dumps(schema)},
    }


async def test_extract_json_is_the_default_and_unchanged(
    client: httpx.AsyncClient, pipeline: FakePipeline
) -> None:
    pipeline.data = INVOICE_DATA
    for kwargs in (
        {},
        {"params": {"format": "json"}},
        {"headers": {**AUTH, "Accept": "text/csv"}, "params": {"format": "json"}},
    ):
        response = await client.post(
            "/api/v1/extract", **{"headers": AUTH, **kwargs}, **upload(TWO_LISTS)
        )
        assert response.headers["content-type"] == "application/json"
        assert set(response.json()) == {"data", "fields", "warnings", "usage"}
        assert "content-disposition" not in response.headers


async def test_extract_csv_by_query_and_by_accept(
    client: httpx.AsyncClient, pipeline: FakePipeline
) -> None:
    pipeline.data = {"numer": "FV/1", "pozycje": POZYCJE, "kwota": 312.45, "tagi": ["a"]}

    by_query = await client.post(
        "/api/v1/extract", headers=AUTH, params={"format": "csv"}, **upload(ONE_LIST)
    )
    by_accept = await client.post(
        "/api/v1/extract", headers={**AUTH, "Accept": "text/csv"}, **upload(ONE_LIST)
    )

    for response in (by_query, by_accept):
        assert response.status_code == 200
        assert response.headers["content-type"] == "text/csv; charset=utf-8"
        assert response.headers["content-disposition"] == 'attachment; filename="extraction.csv"'
        assert "x-schemalens-flattening" not in response.headers
        assert rows(response.content)[1] == ["FV/1", "Paliwo", "2", "150.0", "312.45", "a"]
        assert "ocr;dur=120" in response.headers["server-timing"]


async def test_extract_csv_reports_partial_flattening_and_include_meta(
    client: httpx.AsyncClient, pipeline: FakePipeline
) -> None:
    pipeline.data = INVOICE_DATA

    response = await client.post(
        "/api/v1/extract",
        headers=AUTH,
        params={"format": "csv", "include_meta": "true"},
        **upload(TWO_LISTS),
    )

    assert response.headers["x-schemalens-flattening"] == "partial"
    assert rows(response.content)[0] == [
        "numer",
        "numer__confidence",
        "pozycje.nazwa",
        "pozycje.ilosc",
        "pozycje.cena",
        "platnosci",
    ]


async def test_extract_xlsx_can_be_read_back(
    client: httpx.AsyncClient, pipeline: FakePipeline
) -> None:
    pipeline.data = INVOICE_DATA

    response = await client.post(
        "/api/v1/extract", headers={**AUTH, "Accept": XLSX}, **upload(TWO_LISTS)
    )

    assert response.status_code == 200 and response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == 'attachment; filename="extraction.xlsx"'
    book = workbook(response.content)
    assert book.sheetnames == ["data", "pozycje", "platnosci", "meta"]
    assert book["pozycje"]["C2"].value == 150.0


async def test_csv_and_xlsx_work_with_presets(client: httpx.AsyncClient) -> None:
    files = {"files": {"file": ("doc.png", png(), "image/png")}, "data": {"preset": "invoice"}}

    as_csv = await client.post("/api/v1/extract", headers=AUTH, params={"format": "csv"}, **files)
    as_xlsx = await client.post("/api/v1/extract", headers=AUTH, params={"format": "xlsx"}, **files)

    header = rows(as_csv.content)[0]
    assert header[:2] == ["invoice_number", "issue_date"]
    assert header[-3:] == ["items.name", "items.quantity", "items.net_value"]
    assert workbook(as_xlsx.content).sheetnames == ["data", "items", "meta"]


@pytest.mark.parametrize(
    "kwargs", [{"params": {"format": "pdf"}}, {"headers": {**AUTH, "Accept": "application/xml"}}]
)
async def test_406_for_unknown_format_before_any_work(
    client: httpx.AsyncClient, pipeline: FakePipeline, kwargs: dict
) -> None:
    response = await client.post("/api/v1/extract", **{"headers": AUTH, **kwargs}, **upload(FLAT))

    assert response.status_code == 406
    assert response.headers["content-type"] == "application/problem+json"
    assert response.json()["type"].endswith("#error-not-acceptable")
    assert pipeline.calls == []


async def test_openapi_documents_formats(app) -> None:
    schema = app.openapi()
    for path, method in (("/api/v1/extract", "post"), ("/api/v1/render", "post")):
        operation = schema["paths"][path][method]
        params = {p["name"]: p for p in operation["parameters"]}
        assert params["format"]["schema"]["enum"] == ["json", "csv", "xlsx"]
        assert "include_meta" in params
        assert set(operation["responses"]["200"]["content"]) == {
            "application/json",
            "text/csv",
            XLSX,
        }
        assert "406" in operation["responses"]
