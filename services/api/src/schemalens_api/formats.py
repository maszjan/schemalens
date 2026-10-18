"""Result formats: JSON (the full response), CSV and XLSX (the extracted data, flattened).

The layout follows the schema, not the data, so the columns of a document type are stable:
an empty list still has its columns, a missing value is an empty cell.
"""

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal

from fastapi import Response

from schemalens_api.problems import ApiError
from schemalens_core.schemas import type_of

Format = Literal["json", "csv", "xlsx"]

MEDIA_TYPES: dict[Format, str] = {
    "json": "application/json",
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
FLATTENING_HEADER = "X-SchemaLens-Flattening"
_LIST_SEPARATOR = " | "
_SHEET_NAME_LIMIT = 31
# A cell starting with one of these is run as a formula by spreadsheet programs.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

Node = dict[str, Any]


def negotiate(format_param: str | None, accept: str | None) -> Format:
    """`?format=` wins; otherwise the first supported type in `Accept`; JSON by default."""
    if format_param is not None:
        if format_param in MEDIA_TYPES:
            return format_param  # type: ignore[return-value]
        raise _not_acceptable(f"Unknown format '{format_param[:20]}'.")
    offered = [part.split(";")[0].strip().lower() for part in (accept or "").split(",")]
    offered = [media for media in offered if media]
    if not offered:
        return "json"
    for media in offered:
        if media in ("application/json", "application/*", "*/*"):
            return "json"
        if media in ("text/csv", "text/*"):
            return "csv"
        if media == MEDIA_TYPES["xlsx"]:
            return "xlsx"
    raise _not_acceptable("None of the types in Accept can be produced.")


def _not_acceptable(detail: str) -> ApiError:
    supported = "format=json|csv|xlsx, or Accept: " + ", ".join(MEDIA_TYPES.values())
    return ApiError(406, "not-acceptable", "Not acceptable", f"{detail} Supported: {supported}.")


# --- flattening --------------------------------------------------------------------------


@dataclass
class _Table:
    """Columns in schema order with the schema node of each (for typing cells)."""

    columns: dict[str, Node] = field(default_factory=dict)
    values: dict[str, Any] = field(default_factory=dict)

    def put(self, path: str, value: Any, node: Node) -> None:
        self.columns[path] = node
        self.values[path] = value


@dataclass
class _ObjectList:
    path: str
    node: Node  # schema of one item
    items: list[Any]


_TEXT: Node = {"type": "string"}


def _is_object_list(node: Node) -> bool:
    return type_of(node) == "array" and type_of(node["items"]) == "object"


def _join(values: Any) -> str | None:
    if not isinstance(values, list):
        return None
    return _LIST_SEPARATOR.join("" if v is None else str(v) for v in values)


def _as_json(value: Any) -> str | None:
    return None if value is None else json.dumps(value, ensure_ascii=False)


def _flatten(value: Any, node: Node, path: str, table: _Table, lists: list[_ObjectList]) -> None:
    """Put scalars (and objects, as dotted columns) into `table`; collect lists of objects."""
    kind = type_of(node)
    if kind == "object":
        inner = value if isinstance(value, dict) else {}
        for key, child in node["properties"].items():
            _flatten(inner.get(key), child, f"{path}.{key}", table, lists)
    elif _is_object_list(node):
        lists.append(_ObjectList(path, node["items"], value if isinstance(value, list) else []))
    elif kind == "array":
        table.put(path, _join(value), _TEXT)
    else:
        table.put(path, value, node)


def _item_table(item: Any, node: Node, prefix: str) -> tuple[_Table, bool]:
    """Columns of one list item. Lists of objects inside an item become JSON text."""
    table = _Table()
    nested: list[_ObjectList] = []
    inner = item if isinstance(item, dict) else {}
    for key, child in node["properties"].items():
        _flatten(inner.get(key), child, f"{prefix}{key}", table, nested)
    for deeper in nested:
        table.put(deeper.path, _as_json(deeper.items or None), _TEXT)
    return table, bool(nested)


# --- CSV ---------------------------------------------------------------------------------


def _csv_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        # Document text must not be able to smuggle a formula into a spreadsheet.
        return "'" + value if value.startswith(_FORMULA_PREFIXES) else value
    return str(value)


def to_csv(body: dict[str, Any], schema: Node, *, include_meta: bool = False) -> tuple[bytes, bool]:
    """CSV of `body["data"]` and whether the flattening was partial.

    The first top-level list of objects is expanded into rows; any further list of
    objects ends up as JSON text in one cell (partial flattening).
    """
    data, meta = body["data"], body["fields"]
    header: list[str] = []
    base: dict[str, Any] = {}
    expanded: _ObjectList | None = None
    expanded_at = 0
    partial = False

    for name, node in schema["properties"].items():
        table = _Table()
        lists: list[_ObjectList] = []
        _flatten(data.get(name), node, name, table, lists)
        header += list(table.columns)
        base.update(table.values)
        if include_meta and type_of(node) not in ("object", "array"):
            header.append(f"{name}__confidence")
            base[f"{name}__confidence"] = meta[name]["confidence"]
        for found in lists:
            if expanded is None and found.path == name:
                expanded, expanded_at = found, len(header)
            else:
                header.append(found.path)
                base[found.path] = _as_json(found.items or None)
                partial = True

    rows = [base]
    if expanded is not None:
        prefix = f"{expanded.path}."
        tables = [_item_table(item, expanded.node, prefix) for item in expanded.items or [None]]
        partial = partial or any(nested for _, nested in tables)
        item_columns = list(tables[0][0].columns)
        header[expanded_at:expanded_at] = item_columns
        rows = [{**base, **table.values} for table, _ in tables]

    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([_csv_cell(row.get(column)) for column in header])
    return ("﻿" + buffer.getvalue()).encode("utf-8"), partial


# --- XLSX --------------------------------------------------------------------------------


def _typed(value: Any, node: Node) -> tuple[Any, str | None]:
    """Cell value and number format: dates become real spreadsheet dates."""
    if isinstance(value, str) and type_of(node) == "string":
        try:
            if node.get("format") == "date":
                return date.fromisoformat(value), "yyyy-mm-dd"
            if node.get("format") == "date-time":
                parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
                return parsed.replace(tzinfo=None), "yyyy-mm-dd hh:mm:ss"
        except ValueError:
            pass
    return value, None


# Look of the workbook: the brand yellow header, quiet borders, banded rows.
_HEADER_FILL = "FFE81A"
_BAND_FILL = "F5F5F4"
_BORDER_COLOR = "D6D3D1"
_CONFIDENCE_FILLS = ((0.5, "D9F2D9"), (0.0, "F8D7DA"))  # as on the page: green, red below 50%
_MIN_WIDTH, _MAX_WIDTH = 10, 60
_FONT = "Calibri"


def _header(path: str, labels: dict[str, str] | None) -> str:
    """Column heading: the raw path, or with `labels` a readable name ("seller / city")."""
    if labels is None:
        return path
    parts = path.split(".")
    named = [
        labels.get(".".join(parts[: i + 1])) or parts[i].replace("_", " ")
        for i in range(len(parts))
    ]
    return " / ".join(part[:1].upper() + part[1:] for part in named)


def _number_format(value: Any, node: Node) -> str | None:
    if isinstance(value, bool):
        return None
    if type_of(node) == "integer":
        return "0"
    if type_of(node) == "number" and isinstance(value, int | float):
        return "#,##0.00"
    return None


def _style(sheet: Any, rows: int, columns: int) -> None:
    """Header, borders, banding, column widths, a frozen header row and filters."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if not columns:
        return
    line = Side(style="thin", color=_BORDER_COLOR)
    border = Border(left=line, right=line, top=line, bottom=line)
    header_fill = PatternFill("solid", start_color=_HEADER_FILL)
    band_fill = PatternFill("solid", start_color=_BAND_FILL)
    widths = [_MIN_WIDTH] * columns
    for row in sheet.iter_rows(min_row=1, max_row=rows + 1, max_col=columns):
        for cell in row:
            column = cell.column - 1
            text = "" if cell.value is None else str(cell.value)
            widths[column] = max(widths[column], min(len(text) + 2, _MAX_WIDTH))
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=len(text) > _MAX_WIDTH)
            if cell.row == 1:
                cell.font = Font(name=_FONT, bold=True, color="1C1917")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            elif cell.row % 2 == 1:
                cell.fill = band_fill
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(columns)}{rows + 1}"
    # Printed or saved as PDF: landscape, as wide as the page.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _write_rows(
    sheet: Any,
    columns: dict[str, Node],
    rows: list[dict[str, Any]],
    labels: dict[str, str] | None = None,
) -> None:
    sheet.append([_header(path, labels) for path in columns])
    for row_index, row in enumerate(rows, start=2):
        for column_index, (path, node) in enumerate(columns.items(), start=1):
            value, number_format = _typed(row.get(path), node)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # never a formula, whatever the text starts with
            number_format = number_format or _number_format(value, node)
            if number_format:
                cell.number_format = number_format
    _style(sheet, len(rows), len(columns))


def _sheet_title(path: str, taken: set[str]) -> str:
    title = re.sub(r"[\[\]:*?/\\]", "_", path)[:_SHEET_NAME_LIMIT] or "list"
    candidate, counter = title, 2
    while candidate.lower() in taken:
        suffix = f"_{counter}"
        candidate = title[: _SHEET_NAME_LIMIT - len(suffix)] + suffix
        counter += 1
    taken.add(candidate.lower())
    return candidate


def to_xlsx(body: dict[str, Any], schema: Node, labels: dict[str, str] | None = None) -> bytes:
    """Workbook: `data` (one row), a sheet per list of objects, and `meta`, styled for reading.

    With `labels` (field path -> name, e.g. in the reader's language) the headings are those
    names; without, the raw field paths, as in CSV. The layout is the same either way.
    """
    from openpyxl import Workbook  # only loaded when a spreadsheet is asked for
    from openpyxl.styles import Font, PatternFill

    data = body["data"]
    table = _Table()
    lists: list[_ObjectList] = []
    for name, node in schema["properties"].items():
        _flatten(data.get(name), node, name, table, lists)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.sheet_properties.tabColor = _HEADER_FILL
    _write_rows(sheet, table.columns, [table.values], labels)

    taken = {"data", "meta"}
    for found in lists:
        tables = [_item_table(item, found.node, "")[0] for item in found.items]
        columns = _item_table(None, found.node, "")[0].columns
        list_sheet = workbook.create_sheet(_sheet_title(found.path, taken))
        item_labels = (
            {
                k[len(found.path) + 1 :]: v
                for k, v in labels.items()
                if k.startswith(found.path + ".")
            }
            if labels is not None
            else None
        )
        _write_rows(list_sheet, columns, [t.values for t in tables], item_labels)

    meta = workbook.create_sheet("meta")
    warnings: list[str] = body["warnings"]
    names = list(schema["properties"])

    def about(name: str) -> list[str]:
        return [w for w in warnings if re.match(rf"{re.escape(name)}(?:[:.\[])", w)]

    meta_columns = dict.fromkeys(("field", "confidence", "evidence", "warnings"), _TEXT)
    meta_rows = [
        {
            "field": labels.get(name, name) if labels else name,
            "confidence": body["fields"][name]["confidence"],
            "evidence": ", ".join(body["fields"][name]["evidence"]),
            "warnings": "; ".join(about(name)),
        }
        for name in names
    ]
    general = [w for w in warnings if not any(w in about(name) for name in names)]
    if general:
        meta_rows.append({"field": "(document)", "warnings": "; ".join(general)})
    _write_rows(meta, meta_columns, meta_rows)
    for index, name in enumerate(names, start=2):
        cell = meta.cell(row=index, column=2)
        cell.number_format = "0%"
        if data.get(name) is None:  # nothing found: no confidence to judge, not a red flag
            cell.font = Font(name=_FONT, color="A8A29E")
            continue
        color = next(fill for floor, fill in _CONFIDENCE_FILLS if cell.value >= floor)
        cell.fill = PatternFill("solid", start_color=color)
    meta.append([])
    meta.append(["usage", "value"])
    usage_header = meta.max_row
    for column in (1, 2):
        meta.cell(row=usage_header, column=column).font = Font(name=_FONT, bold=True)
        meta.cell(row=usage_header, column=column).fill = PatternFill(
            "solid", start_color=_HEADER_FILL
        )
    for key, value in body["usage"].items():
        meta.append([key, value])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- HTTP --------------------------------------------------------------------------------


def render(
    fmt: Format,
    body: dict[str, Any],
    schema: Node,
    *,
    include_meta: bool,
    filename: str,
    labels: dict[str, str] | None = None,
) -> Response:
    """The result as a downloadable file. JSON is returned by the routes themselves."""
    headers = {"Content-Disposition": f'attachment; filename="{filename}.{fmt}"'}
    if fmt == "csv":
        content, partial = to_csv(body, schema, include_meta=include_meta)
        if partial:
            headers[FLATTENING_HEADER] = "partial"
        return Response(content, media_type="text/csv; charset=utf-8", headers=headers)
    return Response(to_xlsx(body, schema, labels), media_type=MEDIA_TYPES["xlsx"], headers=headers)
